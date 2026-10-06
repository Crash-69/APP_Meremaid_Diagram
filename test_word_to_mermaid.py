import io
import json
import unittest
from unittest.mock import patch

from docx import Document

from word_to_mermaid_llm_ultimate import (
    _parse_ascii_flow,
    _valid_mermaid,
    convert_docx,
    convert_docx_fallback,
    extract_flow_section,
    markdown_to_mermaid,
    validate_mermaid_with_ollama,
)


FLOW = """START
 |
 V
Controllo autorizzazioni
 |
 +--> KO --> Fine
 |
 V
Verifica periodi chiusi
 |
 V
Estrazione CR
 |
 V
Visualizzazione ALV
 |
 V
Selezione CR
 |
 V
Rilascio Task
 |
 V
Rilascio CR
 |
 V
FINE
"""


class AsciiFlowTests(unittest.TestCase):
    def test_authorization_branch_keeps_main_path(self):
        labels, edges = _parse_ascii_flow(FLOW)
        names = [label.casefold() for label in labels]
        connections = {(names[source], names[target]) for source, target in edges}
        main_path = ['start', 'controllo autorizzazioni', 'verifica periodi chiusi',
                     'estrazione cr', 'visualizzazione alv', 'selezione cr',
                     'rilascio task', 'rilascio cr', 'fine']
        for connection in zip(main_path, main_path[1:]):
            self.assertIn(connection, connections)
        self.assertIn(('controllo autorizzazioni', 'ko'), connections)
        self.assertIn(('ko', 'fine'), connections)
        self.assertNotIn(('fine', 'verifica periodi chiusi'), connections)
        self.assertEqual(names.count('fine'), 1)

    def test_section_and_fenced_flow(self):
        markdown = '# Introduzione\n- Fuori\n\n# 6. Diagramma di Flusso\n\n```text\n' + FLOW + '```\n\n# Note\n- Fuori\n'
        diagram = markdown_to_mermaid(extract_flow_section(markdown))
        self.assertTrue(_valid_mermaid(diagram))
        self.assertIn('Controllo autorizzazioni', diagram)
        self.assertIn('Rilascio CR', diagram)
        self.assertNotIn('Fuori', diagram)

    def test_indented_flow(self):
        markdown = '# Diagramma di Flusso\n\n' + ''.join('    ' + line + '\n' for line in FLOW.splitlines())
        self.assertIn('Controllo autorizzazioni', markdown_to_mermaid(markdown))

    def test_flow_without_code_fences(self):
        for content in [FLOW, '\n\n'.join(FLOW.splitlines())]:
            diagram = markdown_to_mermaid('# 6. Diagramma di Flusso\n\n' + content)
            self.assertEqual(diagram.count('["'), 12)
            self.assertIn('Rilascio CR', diagram)

    def test_real_word_with_normal_paragraphs(self):
        document = Document()
        document.add_heading('6. Diagramma di Flusso', 1)
        for line in FLOW.splitlines():
            document.add_paragraph(line)
        stream = io.BytesIO()
        document.save(stream)
        content = stream.getvalue()
        markdown, _ = convert_docx(content, 'prova.docx', None)
        for source in [markdown, convert_docx_fallback(content)]:
            diagram = markdown_to_mermaid(extract_flow_section(source))
            self.assertEqual(diagram.count('["'), 12)
            self.assertTrue(_valid_mermaid(diagram))

    def test_ollama_cannot_remove_steps(self):
        original = markdown_to_mermaid('# Diagramma di Flusso\n\n' + FLOW)
        shortened = 'graph TD\n  n0["Documento"]\n  n1["Diagramma di Flusso"]\n  n0 --> n1'
        response = io.BytesIO(json.dumps({'message': {'content': shortened}}).encode())
        with patch('word_to_mermaid_llm_ultimate.urlopen', return_value=response):
            diagram, warning = validate_mermaid_with_ollama(original, 'test')
        self.assertEqual(diagram, original)
        self.assertIsNotNone(warning)

    def test_ollama_can_preserve_graph_with_different_header(self):
        original = markdown_to_mermaid('# Diagramma di Flusso\n\n' + FLOW)
        candidate = original.replace('graph TD', 'flowchart TD', 1)
        response = io.BytesIO(json.dumps({'message': {'content': candidate}}).encode())
        with patch('word_to_mermaid_llm_ultimate.urlopen', return_value=response):
            self.assertEqual(validate_mermaid_with_ollama(original, 'test'), (candidate, None))

    def test_regular_headings_and_lists_unchanged(self):
        self.assertEqual(
            markdown_to_mermaid('# Titolo\n- Primo\n  - Secondo'),
            'graph TD\n  n0["Documento"]\n  n1["Titolo"]\n  n0 --> n1\n'
            '  n2["Primo"]\n  n1 --> n2\n  n3["Secondo"]\n  n2 --> n3',
        )

    def test_prose_and_incomplete_arrows_are_not_flows(self):
        for text in ['testo\naltra riga', 'A\n |\n V',
                     'A\n |\n +--> KO -->\n |\n V\nFine']:
            with self.subTest(text=text):
                self.assertIsNone(_parse_ascii_flow(text))

    def test_repeated_actions_remain_distinct(self):
        labels, _ = _parse_ascii_flow('A\n |\n V\nA\n |\n V\nFine')
        self.assertEqual(labels, ['A', 'A', 'Fine'])


if __name__ == '__main__':
    unittest.main()